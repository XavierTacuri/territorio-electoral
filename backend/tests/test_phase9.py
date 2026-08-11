from datetime import date,timedelta
from pathlib import Path
from tempfile import TemporaryDirectory

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.models.alerts import AlertAcknowledgement,AlertRule,OperationalAlert
from app.models.operational import Commitment
from app.models.reports import ReportArtifact,ReportRun,ReportTemplate
from app.reports.common import sanitize_excel_text
from app.reports.excel_renderer import ExcelRenderer
from app.reports.pdf_renderer import PDFRenderer
from app.schemas.alerts import AlertActionRequest,AlertEvaluationRequest
from app.schemas.campaign import CampaignCreate
from app.schemas.reports import ReportGenerationRequest
from app.scripts.seed_gualaceo import seed as seed_gualaceo
from app.scripts.seed_reports_and_alerts import RULES,TEMPLATE_NAMES,seed
from app.services.alert_service import AlertService
from app.services.campaign_service import CampaignService
from app.services.report_storage_service import LocalReportStorage


@pytest.fixture
def phase9_context(db,admin):
    _,canton,parishes=seed_gualaceo(db);db.commit()
    campaign=CampaignService(db).create(CampaignCreate(name="Fase 9",slug="fase-9",canton_id=canton.id,office_type="MAYOR",election_name="Seccionales 2027",election_date=date(2027,2,14),start_date=date(2026,7,1),status="DRAFT"),admin)
    seed(db);db.commit();return campaign,parishes


def test_seed_reports_and_alerts_is_idempotent(db):
    assert seed(db)==(11,21);db.commit();assert seed(db)==(11,21);db.commit()
    assert len(list(db.scalars(select(ReportTemplate))))==len(TEMPLATE_NAMES)
    assert len(list(db.scalars(select(AlertRule))))==len(RULES)


def test_models_use_functional_dates(db,admin,phase9_context):
    campaign,_=phase9_context;template=db.scalar(select(ReportTemplate))
    run=ReportRun(campaign_id=campaign.id,report_template_id=template.id,requested_format="PDF",status="PENDING",report_date=date(2026,8,3),date_from=date(2026,8,1),date_to=date(2026,8,3),filters={},resolved_scope={},title="Informe",requested_by_user_id=admin.id)
    db.add(run);db.flush();artifact=ReportArtifact(report_run_id=run.id,format="PDF",original_download_name="informe.pdf",storage_key="a.pdf",mime_type="application/pdf",size_bytes=10,sha256="a"*64,expires_on=date(2026,9,2));db.add(artifact);db.commit()
    assert isinstance(run.report_date,date) and isinstance(artifact.expires_on,date)


def test_pdf_is_valid_and_dates_are_visible(tmp_path):
    path=tmp_path/"report.pdf";PDFRenderer(10).render(path,"Informe ejecutivo","Campaña Gualaceo",date(2026,8,3),(date(2026,7,5),date(2026,8,3)),[{"title":"Resumen","headers":["Dato","Valor"],"rows":[["Actividades",8]]}])
    content=path.read_bytes();assert content.startswith(b"%PDF") and len(content)>500


def test_xlsx_is_valid_and_has_no_macros(tmp_path):
    path=tmp_path/"report.xlsx";ExcelRenderer().render(path,"Informe","Campaña",date(2026,8,3),(None,None),[{"title":"Actividades","headers":["Nombre"],"rows":[["Recorrido"]]}])
    wb=load_workbook(path);assert {"Resumen","Actividades"}<=set(wb.sheetnames);assert wb["Actividades"].freeze_panes=="A2";assert wb.vba_archive is None


def test_xlsx_serializes_aggregated_structures_and_neutralizes_formulas(tmp_path):
    path=tmp_path/"aggregates.xlsx"
    ExcelRenderer().render(path,"Informe","Campana",date(2026,8,3),(None,None),[{"title":"Resumen agregado","headers":["Dato","Valor"],"rows":[["Prioridades",{"HIGH":1}],["Valores",[1,2]],["Formula","=1+1"]]}])
    wb=load_workbook(path);sheet=wb["Resumen agregado"]
    assert sheet["B2"].value=="{'HIGH': 1}" and sheet["B3"].value=="[1, 2]"
    assert sheet["B4"].value=="'=1+1" and wb.vba_archive is None


@pytest.mark.parametrize("value",["=1+1","+SUM(A1:A2)","-2+3","@cmd"])
def test_formula_injection_is_neutralized(value):assert sanitize_excel_text(value)=="'"+value


def test_storage_checksum_path_and_expiry(tmp_path):
    storage=LocalReportStorage(str(tmp_path),1);source=tmp_path/"tmp.bin";source.write_bytes(b"safe")
    key,size,digest=storage.store(source,"pdf");assert size==4 and len(digest)==64 and storage.resolve(key).parent==tmp_path.resolve()
    with pytest.raises(ValueError):storage.resolve("../outside.pdf")
    assert storage.delete(key) is True and storage.delete(key) is False


def test_overdue_alert_deduplicates_and_updates(db,admin,phase9_context):
    campaign,parishes=phase9_context
    db.add(Commitment(campaign_id=campaign.id,title="Revisión",priority="HIGH",status="PENDING",due_date=date(2026,7,31),parish_id=parishes[0].id,created_by_user_id=admin.id,is_active=True));db.commit()
    service=AlertService(db);request=AlertEvaluationRequest(rule_codes=["OVERDUE_COMMITMENT"],as_of_date=date(2026,8,3))
    first=service.evaluate(campaign.id,admin,request);second=service.evaluate(campaign.id,admin,request)
    assert first["created"]==1 and second["created"]==0
    alerts=list(db.scalars(select(OperationalAlert)));assert len(alerts)==1 and alerts[0].evidence["days_overdue"]==3


def test_alert_acknowledge_resolve_and_history(db,admin,phase9_context):
    campaign,parishes=phase9_context;rule=db.scalar(select(AlertRule).where(AlertRule.code=="OVERDUE_COMMITMENT"))
    alert=OperationalAlert(alert_rule_id=rule.id,campaign_id=campaign.id,severity="WARNING",status="OPEN",title="Compromiso vencido",message="Mensaje neutral.",detected_date=date(2026,8,3),last_seen_date=date(2026,8,3),parish_id=parishes[0].id,fingerprint="f"*64,evidence={},is_active=True);db.add(alert);db.commit()
    service=AlertService(db);service.action(campaign.id,alert.id,"ACKNOWLEDGE",AlertActionRequest(action_date=date(2026,8,3),note="Revisión técnica."),admin);service.action(campaign.id,alert.id,"RESOLVE",AlertActionRequest(action_date=date(2026,8,3)),admin)
    assert alert.status=="RESOLVED" and alert.resolved_date==date(2026,8,3);assert len(list(db.scalars(select(AlertAcknowledgement))))==2


def test_report_and_alert_http_require_auth(client,phase9_context):
    campaign,_=phase9_context
    assert client.get("/api/v1/report-templates").status_code==401
    assert client.get(f"/api/v1/campaigns/{campaign.id}/reports").status_code==401
    assert client.get(f"/api/v1/campaigns/{campaign.id}/alerts").status_code==401


def test_template_http_and_alert_evaluation(client,admin_headers,phase9_context):
    campaign,_=phase9_context
    templates=client.get("/api/v1/report-templates",headers=admin_headers);assert templates.status_code==200 and len(templates.json())==11
    assert "PARISH_TERRITORIAL_PROFILE" in {item["code"] for item in templates.json()}
    evaluated=client.post(f"/api/v1/campaigns/{campaign.id}/alerts/evaluate",headers=admin_headers,json={"rule_codes":["OVERDUE_COMMITMENT"],"as_of_date":"2026-08-03"});assert evaluated.status_code==200
    body=evaluated.text;assert "created_at" not in body and "updated_at" not in body
