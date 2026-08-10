from datetime import date

from openpyxl import load_workbook
from pypdf import PdfReader

from app.reports.current_election import build_current_election_sections
from app.reports.excel_renderer import ExcelRenderer
from app.reports.pdf_renderer import PDFRenderer


def report_data():
    names = ["Gualaceo", "Córdova", "Jadán", "Simón Bolívar", "Daniel Córdova", "Luis Cordero", "Mariano Moreno", "Remigio Crespo", "Zhidmad"]
    parishes = []
    for index, name in enumerate(names):
        current = 3800 + index
        parishes.append({"name": name, "registered_voters_current": current, "historical_2019": {"turnout_rate": .6858}, "historical_2023": {"turnout_rate": .7189}, "projection": {"low": .6808, "central": .71, "high": .7307, "expected_voters_central": round(current * .71)}, "data_quality_status": "MEDIUM", "demographics": {"POP_TOTAL": 4800}})
    return {"snapshot": {"snapshot_date": "2026-07-16", "registered_voters": 34784, "male_voters": 15432, "female_voters": 19352, "juntas": 112}, "historical": {"2019": {"registered_voters": 44019, "ballots_cast": 30190, "turnout_rate": .6858}, "2023": {"registered_voters": 38406, "ballots_cast": 27610, "turnout_rate": .7189}}, "projection": {"model_code": "TURNOUT_HISTORICAL_WEIGHTED_V1", "model_version": "1.0", "expected_voters_low": 23680, "expected_voters_central": 24697, "expected_voters_high": 25416}, "warnings": [{"code": "REGISTRATION_SERIES_BREAK"}], "parishes": parishes, "report_context": {"canton_name": "Gualaceo", "election_name": "Elecciones Seccionales y CPCCS 2027", "sources": [{"institution": "Institución Estadística", "dataset": "Contexto Demográfico", "reference_year": 2022}]}}


def test_pdf_xlsx_unicode_parity_and_regression(tmp_path):
    sections = build_current_election_sections(report_data())
    pdf, xlsx = tmp_path / "informe.pdf", tmp_path / "informe.xlsx"
    PDFRenderer().render(pdf, "Informe ejecutivo — Elección actual", "Gualaceo2026", date(2026, 8, 9), (None, None), sections)
    ExcelRenderer().render(xlsx, "Informe ejecutivo — Elección actual", "Gualaceo2026", date(2026, 8, 9), (None, None), sections)
    pdf_text = "\n".join(page.extract_text() or "" for page in PdfReader(pdf).pages)
    wb = load_workbook(xlsx, data_only=True)
    xlsx_text = "\n".join(wb.sheetnames) + "\n" + "\n".join(str(cell.value or "") for ws in wb for row in ws.iter_rows() for cell in row)
    words = ["Elección", "Participación", "Proyección", "Versión", "Metodología", "Institución", "Estadística", "Demográfico", "Córdova", "Jadán", "Simón Bolívar"]
    for text in (pdf_text, xlsx_text):
        assert all(word in text for word in words), [word for word in words if word not in text]
        assert not any(marker in text for marker in ("Ã", "Â", "�"))
        assert "TURNOUT_HISTORICAL_WEIGHTED_V1" in text and "1.0" in text
    assert all(value in pdf_text for value in ("34.784", "23.680", "24.697", "25.416"))
    assert all(value in xlsx_text for value in ("34784", "23680", "24697", "25416"))
    assert len(report_data()["parishes"]) == 9
    historical = wb["Participación histórica"]
    participation_cells = [cell for row in historical.iter_rows() for cell in row if cell.value == .6858]
    assert participation_cells and participation_cells[0].number_format == "0.00%"
    assert wb["Resumen"]["B3"].value.date() == date(2026, 8, 9)
