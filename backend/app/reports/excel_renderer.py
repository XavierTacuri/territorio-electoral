from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.reports.common import sanitize_excel_text


class ExcelRenderer:
    def render(self, path: Path, title: str, campaign_name: str, report_date: date, period: tuple[date | None, date | None], sections: list[dict[str, Any]]):
        wb = Workbook(); summary_ws = wb.active; summary_ws.title = "Resumen"
        summary_ws.append([title]); summary_ws.append(["Campaña", campaign_name]); summary_ws.append(["Fecha del informe", report_date]); summary_ws["B3"].number_format = "DD/MM/YYYY"
        for section in sections:
            ws = summary_ws if section["title"] == "Resumen" else wb.create_sheet(str(section["title"])[:31])
            headers, formats = section.get("headers", []), section.get("formats", [])
            if section["title"] == "Resumen":
                if section.get("text"): ws.append([section["text"]])
            elif section.get("subtitle") or section.get("text"):
                ws.append([section["title"]])
                if section.get("subtitle"): ws.append([section["subtitle"]])
                if section.get("text"): ws.append([section["text"]])
            if headers:
                header_row = ws.max_row + 1 if ws.max_row > 1 or ws["A1"].value is not None else 1
                ws.append(headers) if header_row > 1 else [setattr(ws.cell(1, i + 1), "value", value) for i, value in enumerate(headers)]
            else: header_row = None
            for row in section.get("rows", []): ws.append([sanitize_excel_text(str(v) if isinstance(v, (dict, list)) else v) for v in row])
            if header_row:
                for cell in ws[header_row]: cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="274C77"); cell.alignment = Alignment(wrap_text=True, vertical="top")
                ws.freeze_panes = f"A{header_row + 1}"; ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ws.max_column)}{ws.max_row}"
                for row in ws.iter_rows(min_row=header_row + 1):
                    for i, cell in enumerate(row):
                        kind = formats[i] if i < len(formats) else None
                        if isinstance(cell.value, date): cell.number_format = "DD/MM/YYYY"
                        elif kind == "integer": cell.number_format = "#,##0"
                        elif kind == "percent": cell.number_format = "0.00%"
                        elif kind in {"decimal", "density"}: cell.number_format = "#,##0.00"
                        elif kind == "date": cell.number_format = "DD/MM/YYYY"
            ws["A1"].font = Font(bold=True, size=14, color="16324F")
            for row in ws.iter_rows():
                for cell in row: cell.alignment = Alignment(wrap_text=True, vertical="top")
            for col in range(1, ws.max_column + 1):
                width = max((len(str(ws.cell(r, col).value or "")) for r in range(1, min(ws.max_row, 80) + 1)), default=12) + 2
                ws.column_dimensions[get_column_letter(col)].width = min(45, max(13, width))
        wb.save(path)
