from datetime import date
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select

from app.api.routes.participation import current_election_analysis
from app.db.session import SessionLocal
from app.models.campaign import Campaign
from app.models.historical import DataSource, DemographicObservation, ElectoralProcess, ElectoralRollSnapshot, ParticipationProjectionRun
from app.models.territory import Canton
from app.models.user import User
from app.reports.current_election import build_current_election_sections
from app.reports.excel_renderer import ExcelRenderer
from app.reports.pdf_renderer import PDFRenderer


def main():
    output = Path("quality-artifacts"); output.mkdir(exist_ok=True)
    with SessionLocal() as db:
        campaign = db.scalar(select(Campaign).order_by(Campaign.created_at))
        user = db.scalar(select(User).where(User.is_superuser.is_(True)))
        analysis = current_election_analysis(campaign.id, user, db)
        run = db.scalar(select(ParticipationProjectionRun).where(ParticipationProjectionRun.campaign_id == campaign.id).order_by(ParticipationProjectionRun.created_at.desc()))
        snapshot = db.get(ElectoralRollSnapshot, run.snapshot_id)
        process = db.get(ElectoralProcess, run.electoral_process_id)
        historical = list(db.scalars(select(ElectoralProcess).where(ElectoralProcess.id.in_(run.historical_process_ids))))
        source_ids = {snapshot.source_id, *(p.source_id for p in historical), *db.scalars(select(DemographicObservation.source_id).where(DemographicObservation.is_official.is_(True)).distinct())}
        sources = list(db.scalars(select(DataSource).where(DataSource.id.in_(source_ids))))
        analysis["report_context"] = {"campaign_name": campaign.name, "canton_name": db.get(Canton, campaign.canton_id).name, "election_name": campaign.election_name,
            "sources": [{"institution": s.institution, "dataset": s.dataset_name, "reference_date": s.reference_date, "reference_year": s.reference_year, "publication_date": s.publication_date, "official_url": s.official_url} for s in sources]}
        sections = build_current_election_sections(analysis)
        stamp = "09-08-2026"
        pdf = output / f"informe-ejecutivo-eleccion-actual-{stamp}.pdf"
        xlsx = output / f"informe-ejecutivo-eleccion-actual-{stamp}.xlsx"
        PDFRenderer().render(pdf, "Informe ejecutivo — Elección actual", campaign.name, date(2026, 8, 9), (None, None), sections)
        ExcelRenderer().render(xlsx, "Informe ejecutivo — Elección actual", campaign.name, date(2026, 8, 9), (None, None), sections)
        wb = load_workbook(xlsx, data_only=True)
        expected_sheets = ["Resumen", "Participación histórica", "Proyección parroquial", "Registro electoral", "Contexto INEC", "Metodología", "Fuentes"]
        assert all(name in wb.sheetnames for name in expected_sheets)
        assert len(analysis["parishes"]) == 9
        regression = (analysis["snapshot"]["registered_voters"], analysis["projection"]["expected_voters_low"], analysis["projection"]["expected_voters_central"], analysis["projection"]["expected_voters_high"])
        assert regression == (34784, 23680, 24697, 25416), regression
        assert analysis["projection"]["model_code"] == "TURNOUT_HISTORICAL_WEIGHTED_V1" and analysis["projection"]["model_version"] == "1.0"
        print({"pdf": str(pdf), "xlsx": str(xlsx), "sheets": wb.sheetnames, "parishes": len(analysis["parishes"]), "regression": regression})


if __name__ == "__main__": main()
